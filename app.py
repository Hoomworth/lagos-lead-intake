import os
import datetime
import re
from functools import wraps
import csv
import io

from dotenv import load_dotenv
load_dotenv()

from flask import Flask, render_template, request, redirect, url_for, flash, session
from flask_sqlalchemy import SQLAlchemy
from flask_migrate import Migrate
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix
from sqlalchemy import text, inspect
from itsdangerous import URLSafeTimedSerializer
from flask_mail import Mail, Message
from openai import OpenAI

app = Flask(__name__)

app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1)

app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'super-secret-key-123')
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get('DATABASE_URL', 'sqlite:///app.db')
if app.config['SQLALCHEMY_DATABASE_URI'] and app.config['SQLALCHEMY_DATABASE_URI'].startswith("postgres://"):
    app.config['SQLALCHEMY_DATABASE_URI'] = app.config['SQLALCHEMY_DATABASE_URI'].replace("postgres://", "postgresql://", 1)
app.config['PERMANENT_SESSION_LIFETIME'] = datetime.timedelta(days=365)
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SESSION_COOKIE_SECURE'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['TEMPLATES_AUTO_RELOAD'] = False
app.config['REMEMBER_COOKIE_DURATION'] = datetime.timedelta(days=365)

# Email Configuration (100% Free via Gmail)
app.config['MAIL_SERVER'] = 'smtp.gmail.com'
app.config['MAIL_PORT'] = 587
app.config['MAIL_USE_TLS'] = True
app.config['MAIL_USERNAME'] = os.environ.get('MAIL_USERNAME')
app.config['MAIL_PASSWORD'] = os.environ.get('MAIL_PASSWORD')
app.config['MAIL_DEFAULT_SENDER'] = os.environ.get('MAIL_USERNAME')

api_key = os.environ.get("OPENAI_API_KEY")

if not api_key:
    print("ERROR: OPENAI_API_KEY not set")
    api_key = "dummy-key-to-prevent-crash-during-migrations"

client = OpenAI(api_key=api_key)

db = SQLAlchemy(app)
migrate = Migrate(app, db)
mail = Mail(app)

# Serializer for generating secure password reset tokens
s = URLSafeTimedSerializer(app.config['SECRET_KEY'])

# -----------------------------
# Database Models
# -----------------------------
class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    full_name = db.Column(db.String(100), nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.datetime.utcnow)
    credits = db.Column(db.Integer, default=5)
    gender = db.Column(db.String(20), nullable=True)
    phone = db.Column(db.String(20), nullable=True)
    company_name = db.Column(db.String(150), nullable=True)
    is_admin = db.Column(db.Boolean, default=False)

    leads = db.relationship('Lead', backref='owner', lazy=True)


class Lead(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    agent_name = db.Column(db.String(100), nullable=False)
    name = db.Column(db.String(100), nullable=False)
    phone = db.Column(db.String(20), nullable=False)
    budget = db.Column(db.String(50), nullable=False)
    location = db.Column(db.String(100), nullable=False)
    property_type = db.Column(db.String(50), nullable=False)
    timeline = db.Column(db.String(50), nullable=False)
    notes = db.Column(db.Text, nullable=True)
    status = db.Column(db.String(20), default='New')
    date_added = db.Column(db.DateTime, default=datetime.datetime.utcnow)
    contacted_at = db.Column(db.DateTime, nullable=True)
    closed_at = db.Column(db.DateTime, nullable=True)
    source = db.Column(db.String(50), default='Manual')
    ai_data = db.Column(db.Text, nullable=True)

    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)


# -----------------------------
# Helpers
# -----------------------------
def analyze_lead(lead):
    score = 0

    # Timeline
    timeline = (lead.timeline or "").lower()
    if "urgent" in timeline:
        score += 30
        intent = "Urgent Buyer"
    elif "month" in timeline:
        score += 20
        intent = "Active Buyer"
    else:
        score += 10
        intent = "Exploring"

    # Budget
    try:
        budget_value = int(''.join(filter(str.isdigit, lead.budget)))
        if budget_value > 50000000:
            score += 30
        elif budget_value > 20000000:
            score += 20
        else:
            score += 10
    except:
        score += 10

    # Property Type
    if "land" in (lead.property_type or "").lower():
        score += 10
    else:
        score += 20

    # Final classification
    if score >= 70:
        quality = "Hot"
        action = "Call immediately"
        timing = "Immediate"
    elif score >= 40:
        quality = "Warm"
        action = "Send WhatsApp first"
        timing = "Soon"
    else:
        quality = "Cold"
        action = "Nurture with follow-up message"
        timing = "Later"

    objections = "May worry about price or location fit"

    return {
        "score": score,
        "quality": quality,
        "intent": intent,
        "action": action,
        "timing": timing,
        "risk": objections
    }

import json
def update_lead_ai_data(lead, new_data):
    try:
        data = json.loads(lead.ai_data) if lead.ai_data else {}
    except:
        data = {}
    data.update(new_data)
    lead.ai_data = json.dumps(data)

def get_current_user():
    user_id = session.get('user_id')
    if not user_id:
        return None
    return User.query.filter_by(id=user_id).first()


def login_required(route_function):
    @wraps(route_function)
    def wrapper(*args, **kwargs):
        if 'user_id' not in session:
            flash('Please log in to continue.', 'error')
            return redirect(url_for('login'))
        return route_function(*args, **kwargs)
    return wrapper

def admin_required(route_function):
    @wraps(route_function)
    def wrapper(*args, **kwargs):
        user = get_current_user()
        if not user or not getattr(user, 'is_admin', False):
            flash('Admin access required.', 'error')
            return redirect(url_for('index'))
        return route_function(*args, **kwargs)
    return wrapper


# -----------------------------
# Message Generators
# -----------------------------
def generate_message_1(lead):
    return f"""Hello {lead.name.title()}, thank you for reaching out regarding your interest in a {lead.property_type} in {lead.location} within your specified budget of {lead.budget}. My name is {lead.agent_name}, and I specialize in premium, verified listings in this exact area.

To ensure I don't overwhelm you with the wrong options, I'd love to quickly confirm a few details. Are you primarily buying for personal use or investment, and do you prefer modern builds or properties with room for renovation?

Also, is your budget slightly flexible if we find an off-market property that perfectly matches your vision?

Once you provide these quick details, I will handpick 2 to 3 of our most solid options and send them over for your review. I look forward to helping you secure the best deal. Best regards, {lead.agent_name}"""


def generate_message_2(lead):
    return f"""Hello {lead.name.title()}, I am checking back with you regarding your ongoing search for a {lead.property_type} in the {lead.location} area.

Over the past few days, I have managed to shortlist a few highly attractive, off-market options that strongly match your criteria in terms of value and quality.

However, before I send these exclusive listings over, I want to ensure my understanding of your preferences is still perfectly accurate. Are you still actively searching, and have there been any changes to your timeline or budget?

Let me know if you would prefer a quick WhatsApp presentation with pictures, or a brief 5-minute call to explain why these specific options stand out. Best regards, {lead.agent_name}"""


def generate_call_script(lead):
    return f"""Hello {lead.name}, my name is {lead.agent_name} calling from Hoomworth CRM. I am reaching out regarding the inquiry you made about a {lead.property_type} in {lead.location}. Our agency specializes in premium properties in that exact neighborhood, and I want to ensure we find exactly what you need. Am I catching you at a good time to speak for just two minutes?

To help me filter out the noise, are you purchasing this primarily for personal use or as an investment? I also see your budget is {lead.budget} and your timeline is {lead.timeline}.

If we find a property that completely blows you away but sits just slightly above that budget, is there any flexibility, or is that a hard ceiling? That makes perfect sense. Based on what you’ve shared, I actually have two specific properties in mind that recently became available.

They haven't been heavily marketed yet, and they align beautifully with your criteria. Here is my proposed next step: I am going to compile the details and exact locations of these properties and send them directly to your WhatsApp. Please review them at your convenience, and if one catches your eye, we can schedule a private viewing. Does that sound fair?"""


def generate_objection_default(lead):
    return f"Hi {lead.name.title()}, I completely understand your hesitation regarding the price. However, premium {lead.property_type}s in {lead.location} rarely stay on the market for long due to incredibly high demand.\n\nEven if this sits slightly outside your initial budget of {lead.budget}, the long-term value and rapid appreciation make it a highly profitable deal. Let's jump on a quick 2-minute call so I can break down the true numbers for you."

def generate_inspection_default(lead):
    return f"Hi {lead.name.title()}, I have just arranged exclusive access to a stunning {lead.property_type} in {lead.location} that perfectly matches your exact requirements.\n\nI am setting up private viewings for my top clients this weekend. Would you prefer a physical inspection on Saturday morning, or a virtual video tour so you can lock it down quickly before it's gone?"

def generate_fomo_default(lead):
    return f"Hi {lead.name.title()}, just a quick market update: The {lead.location} area is currently seeing a massive surge in demand. Prices for a {lead.property_type} are projected to jump significantly in the coming months.\n\nIf you are ready to move forward with your {lead.budget} budget, securing a property right now is the smartest investment decision you can make. Let me know if we should lock something down today."

def generate_offmarket_default(lead):
    return f"Hi {lead.name.title()}, an exclusive off-market {lead.property_type} just came up in {lead.location} that perfectly fits your {lead.budget} budget.\n\nIt is not public yet, and I am only showing it to my VIP clients right now. Let me know if I should send the pictures directly to your WhatsApp before I open it up to other buyers."


# -----------------------------
# Auth Routes
# -----------------------------
@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        full_name = request.form.get('full_name')
        email = request.form.get('email').lower()
        password = request.form.get('password')
        confirm = request.form.get('confirm_password')
        gender = request.form.get('gender')
        phone = request.form.get('phone')
        company_name = request.form.get('company_name')

        if not full_name or not email or not password or not confirm or not gender or not phone:
            flash('Fill all fields', 'error')
            return redirect(url_for('register'))

        if password != confirm:
            flash('Passwords do not match', 'error')
            return redirect(url_for('register'))

        # Password strength validation
        if len(password) < 8 or not re.search(r"[a-z]", password) or not re.search(r"[A-Z]", password) or not re.search(r"[0-9]", password):
            flash('Password must be at least 8 characters long and include an uppercase letter, a lowercase letter, and a number.', 'error')
            return redirect(url_for('register'))

        if User.query.filter_by(email=email).first():
            flash('Email already exists', 'error')
            return redirect(url_for('register'))

        user = User(
            full_name=full_name,
            email=email,
            password_hash=generate_password_hash(password),
            gender=gender,
            phone=phone,
            company_name=company_name
        )

        db.session.add(user)
        db.session.commit()

        flash('Account created. Login now.', 'success')
        return redirect(url_for('login'))

    return render_template('register.html')


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email = request.form.get('email').lower()
        password = request.form.get('password')

        user = User.query.filter_by(email=email).first()

        if not user or not check_password_hash(user.password_hash, password):
            flash('Invalid email or password', 'error')
            return redirect(url_for('login'))

        session['user_id'] = user.id
        session['user_name'] = user.full_name
        session.permanent = True

        if user.credits is None:
            user.credits = 5
            db.session.commit()
        

        return redirect(url_for('index'))

    return render_template('login.html')


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


@app.route('/forgot_password', methods=['GET', 'POST'])
def forgot_password():
    if request.method == 'POST':
        email = request.form.get('email').lower()
        user = User.query.filter_by(email=email).first()
        
        if user:
            # Generate a secure token valid for 1 hour
            token = s.dumps(email, salt='email-confirm')
            reset_url = url_for('reset_password', token=token, _external=True)
            
            # For now, print the link directly to the server logs so you can test it
            print("\n" + "="*50, flush=True)
            print(f"PASSWORD RESET LINK FOR {email}:", flush=True)
            print(reset_url, flush=True)
            print("="*50 + "\n", flush=True)
            
            # Actually send the email to the user's inbox
            if app.config['MAIL_USERNAME'] and app.config['MAIL_PASSWORD']:
                try:
                    msg = Message("Reset Your Password - Hoomworth CRM", recipients=[email])
                    msg.body = f"Hello,\n\nTo reset your password, please click the link below:\n\n{reset_url}\n\nIf you did not request a password reset, please ignore this email.\n\nBest regards,\nHoomworth CRM"
                    mail.send(msg)
                    print(f"Successfully sent reset email to {email}", flush=True)
                except Exception as e:
                    print(f"Failed to send email: {e}", flush=True)

            flash('If an account exists, a password reset link has been generated. Check your server logs!', 'success')
        else:
            # Display the same message even if the email doesn't exist (security best practice)
            flash('If an account exists, a password reset link has been generated. Check your server logs!', 'success')
            
        return redirect(url_for('login'))
        
    return render_template('forgot_password.html')


@app.route('/reset_password/<token>', methods=['GET', 'POST'])
def reset_password(token):
    try:
        # Verify the token is valid and hasn't expired (3600 seconds = 1 hour)
        email = s.loads(token, salt='email-confirm', max_age=3600)
    except Exception:
        flash('The reset link is invalid or has expired.', 'error')
        return redirect(url_for('forgot_password'))

    if request.method == 'POST':
        password = request.form.get('password')
        confirm = request.form.get('confirm_password')

        if password != confirm:
            flash('Passwords do not match.', 'error')
            return redirect(url_for('reset_password', token=token))

        user = User.query.filter_by(email=email).first()
        if user:
            user.password_hash = generate_password_hash(password)
            db.session.commit()
            flash('Your password has been successfully updated! You can now log in.', 'success')
            return redirect(url_for('login'))

    return render_template('reset_password.html', token=token)


@app.route('/admin')
@login_required
@admin_required
def admin():
    users = User.query.all()
    return render_template('admin.html', current_user=get_current_user(), users=users)


@app.route('/admin/add_credits/<int:user_id>', methods=['POST'])
@login_required
@admin_required
def admin_add_credits(user_id):
    user = User.query.get_or_404(user_id)
    try:
        credits_to_add = int(request.form.get('credits', 0))
        user.credits += credits_to_add
        db.session.commit()
        flash(f"Successfully added {credits_to_add} credits to {user.full_name}.", "success")
    except ValueError:
        flash("Invalid credit amount.", "error")
    return redirect(url_for('admin'))


@app.route('/profile', methods=['GET', 'POST'])
@login_required
def profile():
    current_user = get_current_user()
    if request.method == 'POST':
        full_name = request.form.get('full_name')
        company_name = request.form.get('company_name')
        new_password = request.form.get('new_password')
        confirm_password = request.form.get('confirm_password')

        if full_name:
            current_user.full_name = full_name
            
        if company_name is not None:
            current_user.company_name = company_name
        
        if new_password:
            if new_password == confirm_password:
                current_user.password_hash = generate_password_hash(new_password)
            else:
                flash("Passwords do not match.", "error")
                return redirect(url_for('profile'))
        
        db.session.commit()
        flash("Profile updated successfully.", "success")
        return redirect(url_for('profile'))
        
    return render_template('profile.html', current_user=current_user)


# -----------------------------
# Lead Routes
# -----------------------------
@app.route('/')
@login_required
def index():
    return render_template('index.html', current_user=get_current_user())


@app.route('/add_lead', methods=['POST'])
@login_required
def add_lead():
    current_user = get_current_user()

    if not current_user:
        return redirect(url_for('login'))

    agent_name = request.form.get('agent_name')
    name = request.form.get('name')
    phone = request.form.get('phone')
    budget = request.form.get('budget')
    location = request.form.get('location')
    property_type = request.form.get('property_type')
    timeline = request.form.get('timeline')
    notes = request.form.get('notes')
    source = request.form.get('source') or 'Manual'

    if not all([agent_name, name, phone, budget, location, property_type, timeline]):
        flash('Fill all fields', 'error')
        return redirect(url_for('leads'))

    lead = Lead(
        agent_name=agent_name,
        name=name,
        phone=phone,
        budget=budget,
        location=location,
        property_type=property_type,
        timeline=timeline,
        notes=notes,
        user_id=current_user.id,
        source=source
    )

    db.session.add(lead)
    db.session.commit()

    return redirect(url_for('set_prospect', lead_id=lead.id))


@app.route('/upload_csv', methods=['POST'])
@login_required
def upload_csv():
    current_user = get_current_user()
    
    if 'file' not in request.files and 'csv_file' not in request.files:
        flash('No file uploaded', 'error')
        return redirect(url_for('leads'))
        
    file = request.files.get('file') or request.files.get('csv_file')
    
    if file.filename == '':
        flash('No file selected', 'error')
        return redirect(url_for('leads'))
        
    if not (file.filename.endswith('.csv') or file.filename.endswith('.xlsx')):
        flash('Please upload a valid .csv or .xlsx file', 'error')
        return redirect(url_for('leads'))

    try:
        leads_added = 0
        if file.filename.endswith('.csv'):
            stream = io.StringIO(file.stream.read().decode("UTF8"), newline=None)
            csv_input = csv.DictReader(stream)
            
            # Standardize headers (make them lowercase and remove spaces)
            csv_input.fieldnames = [header.strip().lower() for header in csv_input.fieldnames]
            
            for row in csv_input:
                if not any(row.values()): continue # Skip completely empty rows
                    
                lead = Lead(
                    agent_name=current_user.full_name,
                    name=row.get('name', 'Unknown Buyer') or 'Unknown Buyer',
                    phone=row.get('phone', 'N/A') or 'N/A',
                    budget=row.get('budget', 'Flexible') or 'Flexible',
                    location=row.get('location', 'Open') or 'Open',
                    property_type=row.get('property_type', 'Any') or 'Any',
                    timeline=row.get('timeline', 'Flexible') or 'Flexible',
                    notes=row.get('notes', 'Imported via Bulk CSV Upload.') or 'Imported via Bulk CSV Upload.',
                    source='CSV Upload',
                    user_id=current_user.id
                )
                db.session.add(lead)
                leads_added += 1
                
        elif file.filename.endswith('.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(file)
            sheet = wb.active
            
            headers = []
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_idx == 0:
                    headers = [str(cell).strip().lower() if cell else f"col_{i}" for i, cell in enumerate(row)]
                    continue
                
                row_dict = dict(zip(headers, row))
                if not any(row_dict.values()): continue
                
                lead = Lead(
                    agent_name=current_user.full_name,
                    name=str(row_dict.get('name', 'Unknown Buyer') or 'Unknown Buyer'),
                    phone=str(row_dict.get('phone', 'N/A') or 'N/A'),
                    budget=str(row_dict.get('budget', 'Flexible') or 'Flexible'),
                    location=str(row_dict.get('location', 'Open') or 'Open'),
                    property_type=str(row_dict.get('property_type', 'Any') or 'Any'),
                    timeline=str(row_dict.get('timeline', 'Flexible') or 'Flexible'),
                    notes=str(row_dict.get('notes', 'Imported via Bulk XLSX Upload.') or 'Imported via Bulk XLSX Upload.'),
                    source='XLSX Upload',
                    user_id=current_user.id
                )
                db.session.add(lead)
                leads_added += 1
            
        db.session.commit()
        flash(f'Successfully imported {leads_added} leads!', 'success')
        
    except ImportError:
        flash('Missing openpyxl library for Excel files. Run: pip install openpyxl', 'error')
    except Exception as e:
        print("CSV Upload Error:", e)
        flash('Error processing file. Ensure it is a valid CSV or XLSX.', 'error')
        
    return redirect(url_for('leads'))


@app.route('/edit_lead/<int:lead_id>', methods=['GET', 'POST'])
@login_required
def edit_lead(lead_id):
    current_user = get_current_user()
    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    if not lead:
        flash("Lead not found.", "error")
        return redirect(url_for('leads'))

    if request.method == 'POST':
        lead.name = request.form.get('name')
        lead.phone = request.form.get('phone')
        lead.budget = request.form.get('budget')
        lead.location = request.form.get('location')
        lead.property_type = request.form.get('property_type')
        lead.timeline = request.form.get('timeline')
        lead.notes = request.form.get('notes')
        
        db.session.commit()
        flash("Lead updated successfully.", "success")
        return redirect(url_for('prospect'))

    return render_template('edit_lead.html', lead=lead, current_user=current_user)


@app.route('/set_prospect/<int:lead_id>')
@login_required
def set_prospect(lead_id):
    session['last_lead_id'] = lead_id
    return redirect(url_for('prospect'))

@app.route('/prospect')
@login_required
def prospect():
    current_user = get_current_user()

    if not current_user:
        return redirect(url_for('login'))

    lead_id = session.get('last_lead_id')
    if not lead_id:
        last_lead = Lead.query.filter_by(user_id=current_user.id)\
            .order_by(Lead.date_added.desc())\
            .first()
        if not last_lead:
            return redirect(url_for('leads'))
        lead_id = last_lead.id
        session['last_lead_id'] = lead_id

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    if not lead:
        return redirect(url_for('leads'))

    # FORMAT PHONE FOR WHATSAPP
    phone = lead.phone.strip()

    if phone.startswith('0'):
        phone = '234' + phone[1:]

    if phone.startswith('+'):
        phone = phone[1:]

    # Load permanently saved AI messages from the database
    try:
        import json
        saved_ai = json.loads(lead.ai_data) if lead.ai_data else {}
    except Exception:
        saved_ai = {}

    ai_whatsapp = saved_ai.get('ai_whatsapp') or session.pop('ai_whatsapp', None)
    ai_sms = saved_ai.get('ai_sms') or session.pop('ai_sms', None)
    ai_email_subject = saved_ai.get('ai_email_subject') or session.pop('ai_email_subject', None)
    ai_email_body = saved_ai.get('ai_email_body') or session.pop('ai_email_body', None)
    analysis = saved_ai.get('ai_analysis') or session.pop('ai_analysis', None)
    ai_followup = saved_ai.get('ai_followup') or session.pop('ai_followup', None)
    ai_script = saved_ai.get('ai_script') or session.pop('ai_script', None)
    ai_objection = saved_ai.get('ai_objection') or session.pop('ai_objection', None)
    ai_inspection = saved_ai.get('ai_inspection') or session.pop('ai_inspection', None)
    ai_fomo = saved_ai.get('ai_fomo') or session.pop('ai_fomo', None)
    ai_offmarket = saved_ai.get('ai_offmarket') or session.pop('ai_offmarket', None)

    message2 = ai_followup if ai_followup else None
    call_script = ai_script if ai_script else None

    # fallback if AI not used
    if not analysis:
        analysis = analyze_lead(lead)


    return render_template(
        'result.html',
        lead=lead,
        current_user=current_user,
        analysis=analysis,
        message1=ai_whatsapp if ai_whatsapp else generate_message_1(lead),
        is_ai_message1=bool(ai_whatsapp),
        message2=message2,
        call_script=call_script,
        phone=phone,

        sms=ai_sms,
        email_subject=ai_email_subject,
        email_body=ai_email_body,
        
        objection_text=ai_objection if ai_objection else generate_objection_default(lead),
        is_ai_objection=bool(ai_objection),
        
        inspection_text=ai_inspection if ai_inspection else generate_inspection_default(lead),
        is_ai_inspection=bool(ai_inspection),
        
        fomo_text=ai_fomo if ai_fomo else generate_fomo_default(lead),
        is_ai_fomo=bool(ai_fomo),
        
        offmarket_text=ai_offmarket if ai_offmarket else generate_offmarket_default(lead),
        is_ai_offmarket=bool(ai_offmarket)
    )


@app.route('/leads')
@login_required
def leads():
    current_user = get_current_user()

    if not current_user:
        return redirect(url_for('login'))

    # ✅ GET FILTER FROM URL
    status = request.args.get('status')
    search = request.args.get('search')

    # ✅ BASE QUERY
    query = Lead.query.filter_by(user_id=current_user.id)

    # ✅ APPLY STATUS FILTER
    if status:
        query = query.filter_by(status=status)

    # ✅ APPLY SEARCH FILTER
    if search:
        query = query.filter(
            db.or_(Lead.name.ilike(f"%{search}%"), Lead.phone.ilike(f"%{search}%"))
        )

    # ✅ FINAL DATA
    all_leads = query.all()

    # ADD THIS
    leads_with_analysis = []

    for lead in all_leads:
        try:
            analysis = analyze_lead(lead)
        except Exception:
            analysis = {"quality": "Cold", "intent": "Exploring", "score": 10, "action": "Follow up", "risk": "N/A"}
        
        # Safe Date Formatting
        try:
            if hasattr(lead.date_added, 'strftime'):
                formatted_date = lead.date_added.strftime('%d %b %Y')
            else:
                formatted_date = str(lead.date_added).split(' ')[0] if lead.date_added else 'N/A'
        except Exception:
            formatted_date = 'N/A'
            
        # Safe WhatsApp Phone Generator
        try:
            wa_phone = str(lead.phone).strip()
            if wa_phone.startswith('0'):
                wa_phone = '234' + wa_phone[1:]
            if wa_phone.startswith('+'):
                wa_phone = wa_phone[1:]
        except Exception:
            wa_phone = ''

        # Safe Scraper Notes & URL Extractor
        extracted_text = ''
        extracted_url = ''
        try:
            if lead.notes and '[LINK]' in lead.notes:
                parts = lead.notes.split('[LINK]')
                extracted_text = parts[0]
                if '[/LINK]' in parts[1]:
                    extracted_url = parts[1].split('[/LINK]')[0].strip()
            else:
                extracted_text = str(lead.notes) if lead.notes else ''
        except Exception:
            extracted_text = str(lead.notes) if lead.notes else ''

        leads_with_analysis.append({
            "lead": lead,
            "analysis": analysis,
            "formatted_date": formatted_date,
            "wa_phone": wa_phone,
            "extracted_text": extracted_text,
            "extracted_url": extracted_url
        })

    # Sort the leads outside of the loop for correctness and performance
    priority_order = {
            "Hot": 1,
            "Warm": 2,
            "Cold": 3
    }
    try:
        leads_with_analysis.sort(
            key=lambda x: (
                priority_order.get(x["analysis"]["quality"], 4),
                -int(x["lead"].id or 0)
            )
        )
    except Exception as e:
        print("Sorting error safely bypassed:", e)

    # ✅ COUNTS
    total_leads = Lead.query.filter_by(user_id=current_user.id).count()
    new_leads = Lead.query.filter_by(user_id=current_user.id, status='New').count()
    contacted_leads = Lead.query.filter_by(user_id=current_user.id, status='Contacted').count()
    closed_leads = Lead.query.filter_by(user_id=current_user.id, status='Closed').count()

    print("USER CREDITS:", current_user.credits)

    return render_template(
        'leads.html',
        leads=leads_with_analysis,
        current_user=current_user,
        total_leads=total_leads,
        new_leads=new_leads,
        contacted_leads=contacted_leads,
        closed_leads=closed_leads,
        credits=current_user.credits,
        search_query=search or ''
    )

# ✅ DELETE ROUTE (FIXED)
@app.route('/delete_lead/<int:lead_id>', methods=['POST'])
@login_required
def delete_lead(lead_id):
    current_user = get_current_user()

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    if not lead:
        return redirect(url_for('leads'))

    db.session.delete(lead)
    db.session.commit()

    flash('Deleted successfully', 'success')
    return redirect(url_for('leads'))


@app.route('/mark_closed/<int:lead_id>')
@login_required
def mark_closed(lead_id):
    current_user = get_current_user()

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    if lead:
        lead.status = 'Closed'
        if not lead.closed_at:
            lead.closed_at = datetime.datetime.utcnow()
        db.session.commit()

    return redirect(url_for('leads'))    


@app.route('/mark_contacted/<int:lead_id>', methods=['POST'])
@login_required
def mark_contacted(lead_id):
    current_user = get_current_user()

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    if lead:
        lead.status = 'Contacted'
        if not lead.contacted_at:
            lead.contacted_at = datetime.datetime.utcnow()
        db.session.commit()

    return '', 200


@app.route('/generate_ai/<int:lead_id>')
@login_required
def generate_ai(lead_id):
    current_user = get_current_user()

    if current_user.credits <= 0:
        flash("No credits left. Please upgrade.", "error")
        return redirect(url_for('prospect'))

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    if not lead:
        return redirect(url_for('leads'))

    prompt = f"""
You are a world-class real estate sales strategist and copywriter, acting as an assistant for an agent named {lead.agent_name}.
Your task is to analyze a client lead and generate a suite of communication materials.

**Core Instructions:**
- **Tone:** The tone must be professional, confident, and highly personable. It should feel like a helpful, expert consultant, not a pushy salesperson.
- **Natural Language:** Use natural, human-like language. Avoid jargon and overly formal phrases. The output should be conversational.
- **No Emojis:** Do NOT use any emojis in your responses.
- **No Placeholders:** Do NOT use placeholders like [Your Name] or [Client Name]. Use the actual names provided in the lead details.
- **Output Format:** You MUST return ONLY valid JSON. Do not include any text or markdown outside of the JSON structure.

**Lead Details to Analyze:**
- Agent's Name: {lead.agent_name}
- Client's Name: {lead.name}
- Desired Location: {lead.location}
- Client's Budget: {lead.budget}
- Property Type: {lead.property_type}
- Client's Timeline: {lead.timeline}

**JSON Structure to Return:**

{{
  "quality": "Analyze the lead to determine if it is 'Hot', 'Warm', or 'Cold'.",
  "intent": "Briefly describe the buyer's likely intent.",
  "score": "Assign a lead score from 1 to 100 based on the quality and intent.",
  "action": "Suggest the single most effective next action for the agent.",
  "timing": "Suggest when to perform the action.",
  "objection": "Predict the most likely objection or concern the client might have.",

  "whatsapp": "Write a detailed, 3-paragraph WhatsApp message. It should build rapport, confirm their core needs, and propose a clear next step. Do NOT use emojis. Separate paragraphs with a double line break (\\n\\n).",
  "sms": "Write a concise and professional SMS message, strictly under 140 characters. Do NOT use emojis.",
  "email_subject": "Create a compelling and professional email subject line. Do NOT use emojis.",
  "email_body": "Write a comprehensive, single-paragraph professional email. It should summarize the opportunity and call to action clearly. Do NOT use emojis."
}}
"""

    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "user", "content": prompt}
            ]
        )

        import json
        ai_data = json.loads(response.choices[0].message.content)

        # Save permanently to Database
        update_lead_ai_data(lead, {
            'ai_whatsapp': ai_data.get("whatsapp"),
            'ai_sms': ai_data.get("sms"),
            'ai_email_subject': ai_data.get("email_subject"),
            'ai_email_body': ai_data.get("email_body"),
            'ai_analysis': {
                "quality": ai_data.get("quality", "Warm"),
                "intent": ai_data.get("intent", "Buyer"),
                "score": ai_data.get("score", 50),
                "action": ai_data.get("action", "Follow up"),
                "timing": ai_data.get("timing", "Soon"),
                "risk": ai_data.get("objection", "No major concern")
            }
        })

        # deduct credit
        current_user.credits -= 1
        db.session.commit()

        flash("AI analysis generated!", "success")

    except Exception as e:
        print("AI ERROR:", e)
        flash("AI failed. Try again.", "error")

    return redirect(url_for('prospect'))


@app.route('/generate_first_contact/<int:lead_id>')
@login_required
def generate_first_contact(lead_id):
    current_user = get_current_user()

    if current_user.credits <= 0:
        flash("No credits left.", "error")
        return redirect(url_for('prospect')) # Stay on prospect

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    prompt = f"Write a highly detailed, 3-paragraph WhatsApp first-contact message to a real estate client named {lead.name} from me, the agent named {lead.agent_name}. They are inquiring about a {lead.property_type} in {lead.location} with a budget of {lead.budget}. Do NOT use any placeholders like [Your Name] or [Recipient's Name]. Do NOT use emojis. Separate each paragraph with a double line break (\\n\\n)."

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )

    update_lead_ai_data(lead, {'ai_whatsapp': response.choices[0].message.content})

    current_user.credits -= 1
    db.session.commit()

    return redirect(url_for('prospect'))


@app.route('/generate_sms/<int:lead_id>')
@login_required
def generate_sms(lead_id):
    current_user = get_current_user()

    if current_user.credits <= 0:
        flash("No credits left.", "error")
        return redirect(url_for('prospect'))

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    prompt = f"Generate a short SMS under 140 characters to {lead.name} from agent {lead.agent_name} regarding a {lead.property_type} in {lead.location}, budget {lead.budget}. Do NOT use placeholders. Do NOT use emojis."

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )

    sms = response.choices[0].message.content

    update_lead_ai_data(lead, {'ai_sms': sms})

    current_user.credits -= 1
    db.session.commit()

    return redirect(url_for('prospect'))


@app.route('/generate_email/<int:lead_id>')
@login_required
def generate_email(lead_id):
    current_user = get_current_user()

    if current_user.credits <= 0:
        flash("No credits left.", "error")
        return redirect(url_for('prospect'))

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    prompt = f"Write a concise but comprehensive, professional real estate email to the client {lead.name}, from me, the agent {lead.agent_name}. It is about a {lead.property_type} in {lead.location}, budget {lead.budget}. Do NOT use any placeholders. Do NOT use emojis. The email should be a single, well-structured paragraph."

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )

    email_body = response.choices[0].message.content

    update_lead_ai_data(lead, {
        'ai_email_subject': f"{lead.property_type} in {lead.location}",
        'ai_email_body': email_body
    })

    current_user.credits -= 1
    db.session.commit()

    return redirect(url_for('prospect'))


@app.route('/generate_followup/<int:lead_id>')
@login_required
def generate_followup(lead_id):
    current_user = get_current_user()

    if current_user.credits <= 0:
        flash("No credits left.", "error")
        return redirect(url_for('prospect'))

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    prompt = f"Write a detailed, 2-paragraph follow-up WhatsApp message to a real estate client named {lead.name} from me, the agent {lead.agent_name}. They showed interest in a {lead.property_type} in {lead.location}. Do NOT use placeholders. Do NOT use emojis. Separate each paragraph with a double line break (\\n\\n)."

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )

    followup = response.choices[0].message.content

    update_lead_ai_data(lead, {'ai_followup': followup})

    current_user.credits -= 1
    db.session.commit()

    return redirect(url_for('prospect'))


@app.route('/generate_script/<int:lead_id>')
@login_required
def generate_script(lead_id):
    current_user = get_current_user()

    if current_user.credits <= 0:
        flash("No credits left.", "error")
        return redirect(url_for('prospect'))

    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()

    prompt = f"""
You are an expert real estate sales coach. Create a conversational phone call script for me, an agent named {lead.agent_name}, calling my client, {lead.name}.

**My Goal:** To build rapport, confirm their needs for a {lead.property_type} in {lead.location}, and schedule a follow-up.

**Instructions for the script:**
1.  **Format:** Structure it as a dialogue. Use "Agent ({lead.agent_name}):" and "Client ({lead.name}):" to show who is speaking.
2.  **Conversational Tone:** The agent's lines should sound natural, confident, and friendly, not robotic.
3.  **Anticipate Client:** After the agent speaks, write a *likely* client response or question. For example, "Client might say: 'I'm busy right now'" or "Client might ask: 'How did you get my number?'".
4.  **Provide Agent's Rebuttal:** After the anticipated client response, provide a perfect, concise rebuttal or answer for the agent.
5.  **Structure:** The script should have a clear opening, a discovery phase (asking 1-2 key questions), and a closing (setting the next step).
6.  **No Placeholders:** Do not use placeholders like [Your Name]. Use the actual names provided.
7.  **No Emojis:** Do NOT use emojis anywhere in the script.
"""

    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )

    script = response.choices[0].message.content

    update_lead_ai_data(lead, {'ai_script': script})

    current_user.credits -= 1
    db.session.commit()

    return redirect(url_for('prospect'))


@app.route('/generate_objection/<int:lead_id>')
@login_required
def generate_objection(lead_id):
    current_user = get_current_user()
    if current_user.credits <= 0:
        flash("No credits left.", "error")
        return redirect(url_for('prospect'))
        
    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()
    prompt = f"Write a highly persuasive, 2-paragraph objection-crusher WhatsApp script addressed to my client {lead.name} from me, the agent {lead.agent_name}. Use this when the client hesitates on a {lead.property_type} in {lead.location} due to their {lead.budget} budget. Give exact word-for-word responses to justify the value and appreciation. Do NOT use placeholders. Do NOT use emojis."
    
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )
    update_lead_ai_data(lead, {'ai_objection': response.choices[0].message.content})
    current_user.credits -= 1
    db.session.commit()
    return redirect(url_for('prospect'))


@app.route('/generate_inspection/<int:lead_id>')
@login_required
def generate_inspection(lead_id):
    current_user = get_current_user()
    if current_user.credits <= 0:
        flash("No credits left.", "error")
        return redirect(url_for('prospect'))
        
    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()
    prompt = f"Write a highly persuasive, 2-paragraph WhatsApp message addressed to my client {lead.name} from me, the agent {lead.agent_name}. Specifically design it to lock in a physical or virtual viewing this weekend for a {lead.property_type} in {lead.location}. Do NOT use placeholders. Do NOT use emojis."
    
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )
    update_lead_ai_data(lead, {'ai_inspection': response.choices[0].message.content})
    current_user.credits -= 1
    db.session.commit()
    return redirect(url_for('prospect'))


@app.route('/generate_fomo/<int:lead_id>')
@login_required
def generate_fomo(lead_id):
    current_user = get_current_user()
    if current_user.credits <= 0:
        flash("No credits left.", "error")
        return redirect(url_for('prospect'))
        
    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()
    prompt = f"Write a short, data-driven 2-paragraph WhatsApp message addressed to my client {lead.name} from me, the agent {lead.agent_name}. Highlight why buying a {lead.property_type} in {lead.location} right now is a smart investment, creating high FOMO (Fear Of Missing Out) for their {lead.budget} budget. Do NOT use placeholders. Do NOT use emojis."
    
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )
    update_lead_ai_data(lead, {'ai_fomo': response.choices[0].message.content})
    current_user.credits -= 1
    db.session.commit()
    return redirect(url_for('prospect'))


@app.route('/generate_offmarket/<int:lead_id>')
@login_required
def generate_offmarket(lead_id):
    current_user = get_current_user()
    if current_user.credits <= 0:
        flash("No credits left.", "error")
        return redirect(url_for('prospect'))
        
    lead = Lead.query.filter_by(id=lead_id, user_id=current_user.id).first()
    prompt = f"Write a short, mysterious 2-paragraph WhatsApp message to {lead.name} from me, the agent {lead.agent_name}, saying an off-market {lead.property_type} just came up in {lead.location} that perfectly matches their budget of {lead.budget}. It’s not public yet, ask if you should send pictures. Do NOT use placeholders. Do NOT use emojis."
    
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )
    update_lead_ai_data(lead, {'ai_offmarket': response.choices[0].message.content})
    current_user.credits -= 1
    db.session.commit()
    return redirect(url_for('prospect'))


@app.route('/insights')
@login_required
def insights():
    current_user = get_current_user()
    if not current_user:
        return redirect(url_for('login'))

    # Fetch all leads for the current user
    leads = Lead.query.filter_by(user_id=current_user.id).all()

    # Calculate status counts
    total_leads = len(leads)
    new_leads = sum(1 for l in leads if l.status == 'New')
    contacted_leads = sum(1 for l in leads if l.status == 'Contacted')
    closed_leads = sum(1 for l in leads if l.status == 'Closed')

    # Calculate quality counts
    hot_leads = warm_leads = cold_leads = 0
    response_times = []
    close_times = []
    for lead in leads:
        analysis = analyze_lead(lead)
        if analysis['quality'] == 'Hot':
            hot_leads += 1
        elif analysis['quality'] == 'Warm':
            warm_leads += 1
        else:
            cold_leads += 1
            
        if isinstance(lead.contacted_at, datetime.datetime) and isinstance(lead.date_added, datetime.datetime):
            diff = (lead.contacted_at - lead.date_added).total_seconds()
            if diff >= 0:
                response_times.append(diff)
                
        if isinstance(lead.closed_at, datetime.datetime) and isinstance(lead.date_added, datetime.datetime):
            diff = (lead.closed_at - lead.date_added).total_seconds()
            if diff >= 0:
                close_times.append(diff)
            
    # Calculate Conversion Rate
    conversion_rate = round((closed_leads / total_leads * 100), 1) if total_leads > 0 else 0

    # Calculate Averages (Response Time & Lead Close Days)
    avg_response_time = 'N/A'
    if response_times:
        avg_sec = sum(response_times) / len(response_times)
        if avg_sec < 3600:
            avg_response_time = f"{int(avg_sec // 60)} mins"
        elif avg_sec < 86400:
            avg_response_time = f"{round(avg_sec / 3600, 1)} hours"
        else:
            avg_response_time = f"{round(avg_sec / 86400, 1)} days"
            
    lead_close_days = 'N/A'
    if close_times:
        avg_sec = sum(close_times) / len(close_times)
        lead_close_days = f"{round(avg_sec / 86400, 1)} days"

    return render_template('insights.html',
                           total_leads=total_leads, new_leads=new_leads,
                           contacted_leads=contacted_leads, closed_leads=closed_leads,
                           hot_leads=hot_leads, warm_leads=warm_leads, cold_leads=cold_leads,
                           conversion_rate=conversion_rate,
                           avg_response_time=avg_response_time,
                           lead_close_days=lead_close_days,
                           current_user=current_user)
   


# -----------------------------
# RUN
# -----------------------------
with app.app_context():
    db.create_all()

    # Auto-upgrade database specifically for Render production
    try:
        inspector = inspect(db.engine)
        if 'lead' in inspector.get_table_names():
            columns = [col['name'] for col in inspector.get_columns('lead')]
            if 'contacted_at' not in columns:
                db.session.execute(text('ALTER TABLE lead ADD COLUMN contacted_at TIMESTAMP'))
                db.session.commit()
                print("Successfully added contacted_at to Render database.")
            if 'closed_at' not in columns:
                db.session.execute(text('ALTER TABLE lead ADD COLUMN closed_at TIMESTAMP'))
                db.session.commit()
                print("Successfully added closed_at to Render database.")
            if 'source' not in columns:
                db.session.execute(text("ALTER TABLE lead ADD COLUMN source VARCHAR(50) DEFAULT 'Manual'"))
                db.session.commit()
                print("Successfully added source to lead table.")
            if 'ai_data' not in columns:
                db.session.execute(text('ALTER TABLE lead ADD COLUMN ai_data TEXT'))
                db.session.commit()
                print("Successfully added ai_data to lead table.")
                
        if 'user' in inspector.get_table_names():
            user_columns = [col['name'] for col in inspector.get_columns('user')]
            if 'gender' not in user_columns:
                db.session.execute(text('ALTER TABLE "user" ADD COLUMN gender VARCHAR(20)'))
                db.session.commit()
                print("Successfully added gender to user table.")
            if 'phone' not in user_columns:
                db.session.execute(text('ALTER TABLE "user" ADD COLUMN phone VARCHAR(20)'))
                db.session.commit()
                print("Successfully added phone to user table.")
            if 'company_name' not in user_columns:
                db.session.execute(text('ALTER TABLE "user" ADD COLUMN company_name VARCHAR(150)'))
                db.session.commit()
                print("Successfully added company_name to user table.")
            if 'is_admin' not in user_columns:
                db.session.execute(text('ALTER TABLE "user" ADD COLUMN is_admin BOOLEAN DEFAULT FALSE'))
                db.session.execute(text('UPDATE "user" SET is_admin = TRUE WHERE id = 1'))
                db.session.commit()
                print("Successfully added is_admin to user table.")

        # Force the very first registered user to ALWAYS be an admin safely after tables exist
        db.session.execute(text('UPDATE "user" SET is_admin = TRUE WHERE id = 1'))
        db.session.commit()
    except Exception as e:
        print(f"Migration check skipped: {e}")
        db.session.rollback()

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 8080)))